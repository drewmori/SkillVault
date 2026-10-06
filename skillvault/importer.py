"""Extract an editable expert-decision draft from completed-work files."""

from __future__ import annotations

from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
import json
import re


@dataclass
class WorkFile:
    name: str
    content_type: str
    data: bytes


@dataclass
class DecisionDraft:
    summary: str
    decision: str
    reasoning: str
    instructions: str
    software: str
    methods: str
    outcome: str
    source_file: str
    goal: str = ""
    chosen_approach: str = ""
    alternatives_considered: str = ""
    constraints: str = ""
    reusable_rule: str = ""
    exceptions: str = ""
    missing_context: list[str] = field(default_factory=list)
    generation_method: str = "Local artifact extraction"
    media: list[dict[str, object]] = field(default_factory=list)
    readable_files: list[str] = field(default_factory=list)
    evidence_preview: str = ""
    warnings: list[str] = field(default_factory=list)
    capture_evidence: dict[str, object] = field(default_factory=dict)


TEXT_SUFFIXES = {
    ".txt", ".md", ".log", ".csv", ".json", ".yaml", ".yml",
    ".toml", ".ini", ".xml", ".rst", ".patch", ".diff", ".ipynb",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".cs", ".go",
    ".rb", ".php", ".sh", ".ps1", ".html", ".css", ".sql", ".eml",
}


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def _email_text(data: bytes) -> str:
    message = BytesParser(policy=policy.default).parsebytes(data)
    parts = [f"Subject: {message.get('subject', '')}"]
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                try:
                    parts.append(part.get_content())
                except (LookupError, UnicodeDecodeError):
                    continue
    else:
        try:
            parts.append(message.get_content())
        except (LookupError, UnicodeDecodeError):
            pass
    return "\n".join(parts)


def _json_text(text: str) -> str:
    """Turn common completed-work JSON fields into extractor-friendly labels."""
    payload = json.loads(text)
    if not isinstance(payload, dict):
        return json.dumps(payload, indent=2)

    aliases = (
        ("Subject", ("subject", "title", "summary", "case", "situation")),
        ("Situation", ("situation", "incident", "problem", "context")),
        ("Decision", ("decision", "resolution", "approved_action")),
        ("Reasoning", ("reasoning", "rationale", "why")),
        ("Software", ("software", "software_and_tools", "systems", "tools")),
        ("Methods", ("methods", "method")),
        ("Outcome", ("outcome", "result", "final_status")),
    )
    labeled: list[str] = []
    used: set[str] = set()
    for label, keys in aliases:
        for key in keys:
            value = payload.get(key)
            if value not in (None, "", [], {}):
                if isinstance(value, list):
                    value = "; ".join(str(item) for item in value)
                elif isinstance(value, dict):
                    value = json.dumps(value, ensure_ascii=False)
                labeled.append(f"{label}: {value}")
                used.add(key)
                break

    for key in ("completed_steps", "steps", "instructions", "actions"):
        value = payload.get(key)
        if value:
            labeled.append("Completed steps:")
            items = value if isinstance(value, list) else [value]
            labeled.extend(f"{index}. {item}" for index, item in enumerate(items, start=1))
            used.add(key)
            break

    remaining = {key: value for key, value in payload.items() if key not in used}
    if remaining:
        labeled.append("Additional evidence:")
        labeled.append(json.dumps(remaining, indent=2, ensure_ascii=False))
    return "\n".join(labeled) or json.dumps(payload, indent=2, ensure_ascii=False)


def _notebook_text(text: str) -> str:
    payload = json.loads(text)
    sections: list[str] = []
    for index, cell in enumerate(payload.get("cells", []), start=1):
        if not isinstance(cell, dict):
            continue
        source = cell.get("source", [])
        source_text = "".join(source) if isinstance(source, list) else str(source or "")
        if source_text.strip():
            sections.append(f"NOTEBOOK {cell.get('cell_type', 'cell').upper()} CELL {index}:\n{source_text.strip()}")
        for output in cell.get("outputs", []) or []:
            if not isinstance(output, dict):
                continue
            output_text = output.get("text", "")
            if isinstance(output_text, list):
                output_text = "".join(output_text)
            if str(output_text).strip():
                sections.append(f"CELL {index} OUTPUT:\n{str(output_text).strip()}")
    return "\n\n".join(sections)


def extract_text(work_file: WorkFile) -> tuple[str, str | None]:
    """Return extracted text and an optional user-facing warning."""
    suffix = Path(work_file.name).suffix.lower()
    if suffix == ".eml":
        return _email_text(work_file.data), None
    if suffix in TEXT_SUFFIXES:
        text = _decode(work_file.data)
        if suffix == ".json":
            try:
                text = _json_text(text)
            except json.JSONDecodeError:
                pass
        elif suffix == ".ipynb":
            try:
                text = _notebook_text(text)
            except (json.JSONDecodeError, AttributeError):
                pass
        return text, None
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(work_file.data))
            return "\n\n".join(f"PAGE {number}:\n{page.extract_text() or ''}" for number, page in enumerate(reader.pages, 1)), None
        except Exception as error:
            return "", f"Could not extract text from {work_file.name}: {error}"
    if suffix == ".pptx":
        try:
            from pptx import Presentation

            presentation = Presentation(BytesIO(work_file.data))
            sections: list[str] = []
            first_title = ""
            explicit_subject = ""
            for slide_number, slide in enumerate(presentation.slides, start=1):
                slide_lines: list[str] = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        slide_lines.append(shape.text.strip())
                        if not first_title:
                            first_title = shape.text.strip().splitlines()[0]
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            slide_lines.append(" | ".join(cell.text.strip() for cell in row.cells))
                notes_frame = getattr(slide.notes_slide, "notes_text_frame", None)
                if notes_frame is not None and notes_frame.text.strip():
                    notes_text = notes_frame.text.strip()
                    if not explicit_subject:
                        subject_match = re.search(r"^Subject\s*:\s*(.+)$", notes_text, re.IGNORECASE | re.MULTILINE)
                        if subject_match:
                            explicit_subject = subject_match.group(1).strip()
                    slide_lines.append(f"SPEAKER NOTES:\n{notes_text}")
                sections.append(f"SLIDE {slide_number}:\n" + "\n".join(slide_lines))
            subject = explicit_subject or first_title or Path(work_file.name).stem.replace("_", " ")
            return f"Subject: {subject}\n" + "\n\n".join(sections), None
        except Exception as error:
            return "", f"Could not extract text from {work_file.name}: {error}"
    if suffix == ".docx":
        try:
            from docx import Document

            document = Document(BytesIO(work_file.data))
            paragraphs: list[str] = []
            numbered_index = 0
            for paragraph_number, paragraph in enumerate(document.paragraphs, 1):
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = (paragraph.style.name or "").lower()
                if "list number" in style_name:
                    numbered_index += 1
                    text = f"{numbered_index}. {text}"
                paragraphs.append(f"PARAGRAPH {paragraph_number}:\n{text}")
            table_rows = []
            for table_number, table in enumerate(document.tables, 1):
                for row_number, row in enumerate(table.rows, 1):
                    table_rows.append(f"TABLE {table_number} ROW {row_number}:")
                    values = [cell.text.strip() for cell in row.cells]
                    if len(values) == 2 and values[0] and values[1]:
                        table_rows.append(f"{values[0]}: {values[1]}")
                    else:
                        table_rows.append(" | ".join(values))
            return "\n".join(paragraphs + table_rows), None
        except Exception as error:
            return "", f"Could not extract text from {work_file.name}: {error}"
    return "", None


def _clean_lines(text: str) -> list[str]:
    return [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]


def _labeled_value(lines: list[str], labels: tuple[str, ...]) -> str:
    pattern = re.compile(rf"^(?:{'|'.join(re.escape(label) for label in labels)})\s*[:\-]\s*(.+)$", re.IGNORECASE)
    label_only = re.compile(rf"^(?:{'|'.join(re.escape(label) for label in labels)})$", re.IGNORECASE)
    for index, line in enumerate(lines):
        match = pattern.match(line)
        if match:
            return match.group(1).strip()
        if label_only.match(line) and index + 1 < len(lines):
            return lines[index + 1].strip()
    return ""


def _first_useful_line(lines: list[str]) -> str:
    ignored = ("source file:", "date:", "from:", "to:", "cc:", "status:", "priority:", "id:", "created:", "updated:")
    for line in lines:
        if len(line) >= 18 and not line.lower().startswith(ignored):
            return re.sub(r"^[#>*\-\d.\s]+", "", line)[:260]
    return "Review the uploaded completed work and capture the expert decision."


def _find_sentence(text: str, keywords: tuple[str, ...]) -> str:
    sentences = re.split(r"(?<=[.!?])\s+|\n+", text)
    for sentence in sentences:
        cleaned = re.sub(r"\s+", " ", sentence).strip(" -*#")
        if len(cleaned) >= 20 and any(re.search(rf"\b{re.escape(keyword)}\b", cleaned, flags=re.IGNORECASE) for keyword in keywords):
            return cleaned[:600]
    return ""


def _extract_steps(lines: list[str], text: str) -> str:
    explicit = []
    for line in lines:
        if re.match(r"^(?:\d+[.)]|[-*•]|step\s+\d+)\s+", line, re.IGNORECASE):
            step = re.sub(r"^(?:\d+[.)]|[-*•]|step\s+\d+)\s+", "", line, flags=re.IGNORECASE).strip()
            if len(step) >= 12:
                explicit.append(step)
    if not explicit:
        action_words = ("checked", "tested", "reviewed", "compared", "captured", "verified", "reproduced", "updated", "escalated", "fixed", "changed", "sent", "rolled back")
        sentences = re.split(r"(?<=[.!?])\s+|\n+", text)
        explicit = [re.sub(r"\s+", " ", sentence).strip(" -*#") for sentence in sentences if len(sentence.strip()) >= 20 and any(word in sentence.lower() for word in action_words)]
    if explicit:
        return "\n".join(f"{index}. {step[:500]}" for index, step in enumerate(explicit[:7], start=1))
    return (
        "1. Review the attached source evidence and confirm the exact situation, scope, and timeline.\n"
        "2. Record the checks and actions that were actually completed.\n"
        "3. Verify the result against the desired outcome.\n"
        "4. Replace these draft steps with the expert's exact process before approval."
    )


def _decision_for(text: str, lines: list[str] | None = None) -> str:
    explicit = _labeled_value(lines or _clean_lines(text), ("decision", "approved decision", "selected approach"))
    for decision in (
        "Follow standard process",
        "Diagnose and verify",
        "Build in small steps",
        "Write and test",
        "Escalate for review",
    ):
        if explicit.lower().startswith(decision.lower()):
            return decision
    lowered = text.lower()
    if any(word in lowered for word in ("powerpoint", "presentation", "slides", "investor", "pitch deck")):
        return "Follow standard process"
    if any(word in lowered for word in ("security", "privacy", "legal", "production access", "delete", "breach", "escalat")):
        return "Escalate for review"
    code_signal = any(word in lowered for word in ("python", "javascript", "typescript", "function", "source code", "pull request", "code change"))
    test_signal = any(word in lowered for word in ("test", "pytest", "unit test", "regression"))
    if code_signal and test_signal:
        return "Write and test"
    if any(word in lowered for word in ("prototype", "pilot", "minimum", "small version", "iterate", "experiment")):
        return "Build in small steps"
    if any(word in lowered for word in ("error", "failed", "failure", "bug", "incident", "logs", "root cause", "reproduce", "discrepancy")):
        return "Diagnose and verify"
    return "Follow standard process"


def _tools_and_methods(text: str) -> tuple[str, str]:
    lines = _clean_lines(text)
    tools = _labeled_value(lines, ("software", "tools", "software and tools"))
    methods = _labeled_value(lines, ("methods", "method"))
    return tools, methods


def draft_from_completed_work(
    files: list[WorkFile],
    employee_notes: str = "",
    decision_context: dict[str, str] | None = None,
) -> DecisionDraft:
    """Build a reviewable draft; never add it to training without approval."""
    extracted_sections: list[str] = []
    readable_files: list[str] = []
    media: list[dict[str, object]] = []
    warnings: list[str] = []

    for work_file in files:
        if work_file.content_type.startswith(("image/", "video/")):
            media.append({"name": work_file.name, "type": work_file.content_type, "bytes": work_file.data})
            continue
        text, warning = extract_text(work_file)
        if warning:
            warnings.append(warning)
        if text.strip():
            readable_files.append(work_file.name)
            extracted_sections.append(f"SOURCE FILE: {work_file.name}\n{text[:30000]}")

    if employee_notes.strip():
        readable_files.append("employee_notes")
        extracted_sections.append(f"EMPLOYEE NOTES:\n{employee_notes.strip()}")
    context = {key: str(value).strip() for key, value in (decision_context or {}).items() if str(value).strip()}
    if context:
        context_labels = {
            "artifact_type": "Artifact type",
            "goal": "Goal",
            "chosen_approach": "Chosen approach",
            "reasoning": "Decision rationale",
            "alternatives_considered": "Alternatives considered",
            "constraints": "Constraints",
            "outcome": "Outcome",
            "reusable_rule": "Reusable rule",
            "exceptions": "Exceptions",
        }
        context_text = "\n".join(f"{context_labels.get(key, key)}: {value}" for key, value in context.items())
        readable_files.append("decision_context")
        extracted_sections.append(f"EXPERT DECISION CONTEXT:\n{context_text}")
    combined = "\n\n".join(extracted_sections)
    lines = _clean_lines(combined)
    summary = _labeled_value(lines, ("issue", "problem", "summary", "subject", "request", "incident", "situation")) or _first_useful_line(lines)
    goal = _labeled_value(lines, ("goal", "objective", "desired outcome", "problem being solved"))
    chosen_approach = _labeled_value(lines, ("chosen approach", "selected approach", "implementation choice", "decision"))
    if not chosen_approach:
        choice_sentence = _find_sentence(combined, ("we chose", "we decided", "decided to", "selected approach", "chose to"))
        chosen_approach = re.split(r"\s+because\s+", choice_sentence, maxsplit=1, flags=re.IGNORECASE)[0].strip()
        chosen_approach = re.sub(r"^(?:reasoning|decision rationale|why)\s*:\s*", "", chosen_approach, flags=re.IGNORECASE)
    reasoning = _labeled_value(lines, ("reasoning", "why", "decision rationale", "root cause")) or _find_sentence(combined, ("because", "root cause", "decided", "chose", "therefore"))
    alternatives_considered = _labeled_value(lines, ("alternatives considered", "alternatives rejected", "other options", "instead of")) or _find_sentence(combined, ("instead of", "rather than", "rejected", "alternative"))
    constraints = _labeled_value(lines, ("constraints", "requirements", "limitations", "tradeoffs")) or _find_sentence(combined, ("constraint", "tradeoff", "limited to", "deadline", "must not"))
    outcome = _labeled_value(lines, ("outcome", "resolution", "result", "final status")) or _find_sentence(combined, ("resolved", "fixed", "restored", "completed", "succeeded", "outcome"))
    # A statement that the outcome is unknown is evidence of a gap, not a result.
    if re.search(
        r"\b(?:no|not|unknown|unmeasured|pending|not yet)\b.{0,60}\b(?:outcome|result|success|impact)\b"
        r"|\b(?:outcome|result|success|impact)\b.{0,60}\b(?:unknown|not measured|not yet|pending|not supplied)\b",
        outcome, flags=re.IGNORECASE,
    ):
        outcome = ""
    reusable_rule = _labeled_value(lines, ("reusable rule", "lesson", "principle", "use this when"))
    exceptions = _labeled_value(lines, ("exceptions", "do not use when", "escalate when", "stop conditions"))
    tools, methods = _tools_and_methods(combined)

    if not combined.strip():
        warnings.append("No readable text was found. Add employee notes or a transcript so SkillVault can draft the decision accurately.")
    if not reasoning:
        reasoning = "The source evidence does not clearly state why the action was chosen. The expert must add the reasoning before approval."
        warnings.append("Decision reasoning was not explicit in the uploaded work.")
    if not outcome:
        outcome = "The final outcome was not explicit in the uploaded work. The expert must record it before approval."
        warnings.append("A final outcome was not found in the uploaded work.")

    missing_context = []
    for label, value in (
        ("the goal or problem being solved", goal),
        ("the exact choice that was made", chosen_approach),
        ("why that choice was better than the alternatives", reasoning if "must add" not in reasoning.lower() else ""),
        ("alternatives that were considered", alternatives_considered),
        ("constraints or tradeoffs", constraints),
        ("the verified outcome", outcome if "must record" not in outcome.lower() else ""),
        ("when this decision should not be reused", exceptions),
    ):
        if not value.strip():
            missing_context.append(label)

    source_names = [work_file.name for work_file in files]
    return DecisionDraft(
        summary=summary,
        decision=_decision_for(combined, lines),
        reasoning=reasoning,
        instructions=_extract_steps(lines, combined),
        software=tools,
        methods=methods,
        outcome=outcome,
        source_file="; ".join(source_names) or "imported_completed_work",
        goal=goal,
        chosen_approach=chosen_approach,
        alternatives_considered=alternatives_considered,
        constraints=constraints,
        reusable_rule=reusable_rule,
        exceptions=exceptions,
        missing_context=missing_context,
        media=media,
        readable_files=readable_files,
        evidence_preview=combined[:16000],
        warnings=warnings,
    )
