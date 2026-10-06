from types import SimpleNamespace
from io import BytesIO
import unittest

from skillvault.source_locations import quote_location
from skillvault.speech import transcribe_recording, TranscriptionError
from skillvault.importer import extract_text, WorkFile


class SourceAndSpeechTests(unittest.TestCase):
    def test_citations_resolve_to_document_units(self):
        for filename, marker in [("deck.pptx", "SLIDE 3"), ("report.pdf", "PAGE 2"), ("memo.docx", "PARAGRAPH 7")]:
            self.assertIn(marker.capitalize(), quote_location(filename, f"{marker}:\nVerified evidence", "Verified evidence"))
        self.assertEqual(quote_location("fix.py", "# comment\nreturn 42", "return 42"), "Line 2")
        with self.assertRaises(ValueError):
            quote_location("fix.py", "return 42", "invented")

    def test_word_paragraph_markers_preserve_decision_fields(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Goal: Speed up checkout")
        document.add_paragraph("Reasoning: Peak traffic increased")
        buffer = BytesIO()
        document.save(buffer)
        text, warning = extract_text(WorkFile("decision.docx", "application/octet-stream", buffer.getvalue()))
        self.assertIsNone(warning)
        self.assertIn("PARAGRAPH 2:\nReasoning: Peak traffic increased", text)

    def test_local_transcript_has_time_references_and_rejects_silence(self):
        class Recognizer:
            def transcribe(self, audio, **kwargs):
                return iter([SimpleNamespace(start=1.2, end=4.5, text=" We chose a smaller rollout. ")]), None
        self.assertEqual(transcribe_recording(b"test", model=Recognizer()), "[1.2s–4.5s] We chose a smaller rollout.")
        with self.assertRaises(TranscriptionError):
            transcribe_recording(b"", model=Recognizer())
        class Silent:
            def transcribe(self, audio, **kwargs):
                return iter([]), None
        with self.assertRaises(TranscriptionError):
            transcribe_recording(b"test", model=Silent())
